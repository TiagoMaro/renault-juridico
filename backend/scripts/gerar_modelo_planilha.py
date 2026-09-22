"""Gera o modelo de planilha que os escritórios devem preencher.

Rode com:  python scripts/gerar_modelo_planilha.py [caminho_de_saida]

O arquivo tem duas abas:
  - "Processos": cabeçalho na linha 1 (como o importador exige) + linhas de exemplo
  - "Como preencher": o que cada coluna aceita, com os nomes alternativos reconhecidos
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AZUL_ESCURO = '0D1B3E'
AZUL_RENAULT = '0035AD'
CINZA_CLARO = 'F1F5F9'
AMARELO = 'FFF7CC'

FONTE = 'Arial'

COLUNAS = [
    # (cabeçalho, largura, obrigatória, formato, o que aceita, nomes alternativos)
    (
        'Status',
        14,
        False,
        None,
        'Ativo ou Inativo. Também aceita: em andamento, aberto, encerrado, arquivado, baixado.',
        'Situação · Situação do Processo',
    ),
    (
        'Autor/Réu',
        34,
        True,
        None,
        'Nome da parte contrária (ou das partes). Linha sem este campo é descartada.',
        'Autor · Réu · Parte · Partes · Autor e Réu',
    ),
    (
        'Número Autos',
        26,
        True,
        None,
        'Número do processo. Com ou sem máscara CNJ — 20 dígitos corridos são formatados automaticamente. '
        'É a chave do processo: reenviar o mesmo número atualiza o registro em vez de duplicar.',
        'Nº Autos · Número do Processo · Processo · Autos · CNJ',
    ),
    (
        'Natureza da ação',
        20,
        False,
        None,
        'Trabalhista, Cível, Consumidor, Tributário, Contratual, Administrativo ou Outros. '
        'Variações são reconhecidas (ex.: "Fiscal" vira Tributário).',
        'Natureza · Tipo de ação · Matéria · Área',
    ),
    ('Vara', 26, False, None, 'Vara ou juízo onde o processo tramita.', 'Juízo · Órgão Julgador'),
    ('Comarca', 22, False, None, 'Comarca / cidade do processo.', 'Cidade · Foro · Localidade'),
    (
        'Data de início',
        16,
        False,
        'DD/MM/YYYY',
        'Data de distribuição/ajuizamento. Aceita dd/mm/aaaa, aaaa-mm-dd e data do Excel.',
        'Data de distribuição · Distribuição · Ajuizamento · Início',
    ),
    (
        'Posição da Renault',
        20,
        False,
        None,
        'ativo ou passivo (vira Polo Ativo / Polo Passivo). Também aceita autora, ré, requerente, requerida.',
        'Posição · Polo · Polo Renault',
    ),
    (
        'Resumo do caso',
        56,
        False,
        None,
        'Descrição curta do objeto da ação. Texto livre.',
        'Resumo · Objeto · Descrição · Síntese',
    ),
    (
        'Defesa',
        12,
        False,
        None,
        'Sim ou Não (defesa apresentada?). Também aceita S/N, realizada, pendente.',
        'Defesa realizada · Contestação',
    ),
    (
        'Fase processual',
        24,
        False,
        None,
        'Conhecimento, Recurso, Execução, Cumprimento de sentença, Encerrado ou Outros.',
        'Fase · Estágio · Fase do Processo',
    ),
    (
        'Movimentações',
        60,
        False,
        None,
        'Andamentos do processo, um por linha (Alt+Enter na mesma célula). '
        'Quando a linha começa com a data, ela vira a data do andamento.',
        'Movimentação · Andamentos · Últimos andamentos',
    ),
    (
        'Valor da causa',
        18,
        False,
        'R$ #,##0.00',
        'Valor atribuído à causa. Aceita 450000, R$ 450.000,00 ou 450,000.00.',
        'Valor Causa · Vlr Causa · Valor',
    ),
    (
        'Valor do risco',
        18,
        False,
        'R$ #,##0.00',
        'Exposição estimada. Define a faixa de risco: Crítico ≥ R$ 1 mi · Alto ≥ R$ 300 mil · '
        'Médio ≥ R$ 50 mil. Processo ativo sem defesa sobe uma faixa.',
        'Valor Risco · Vlr Risco · Risco · Provisão · Exposição · Contingência',
    ),
]

# Dados fictícios — servem só para mostrar o formato esperado de cada coluna.
EXEMPLOS = [
    {
        'Status': 'Ativo',
        'Autor/Réu': 'João Carlos Mendes',
        'Número Autos': '0001234-56.2021.5.09.0015',
        'Natureza da ação': 'Trabalhista',
        'Vara': '3ª Vara do Trabalho',
        'Comarca': 'Curitiba',
        'Data de início': date(2021, 3, 15),
        'Posição da Renault': 'passivo',
        'Resumo do caso': (
            'Ex-funcionário pleiteia verbas rescisórias, horas extras não remuneradas e danos morais '
            'após demissão sem justa causa.'
        ),
        'Defesa': 'Sim',
        'Fase processual': 'Recurso',
        'Movimentações': (
            '15/12/2024 - Despacho: recurso ordinário recebido. Vista à parte contrária.\n'
            '28/11/2024 - Recurso ordinário interposto pela reclamada.\n'
            '10/10/2024 - Sentença proferida. Condenação parcial.'
        ),
        'Valor da causa': 450000,
        'Valor do risco': 185000,
    },
    {
        'Status': 'Ativo',
        'Autor/Réu': 'Maria Aparecida Silva',
        'Número Autos': '0002345-78.2022.5.09.0001',
        'Natureza da ação': 'Trabalhista',
        'Vara': '1ª Vara do Trabalho',
        'Comarca': 'São José dos Pinhais',
        'Data de início': date(2022, 7, 8),
        'Posição da Renault': 'passivo',
        'Resumo do caso': 'Pedido de horas extras, adicional noturno e intervalo intrajornada não concedido.',
        'Defesa': 'Sim',
        'Fase processual': 'Conhecimento',
        'Movimentações': '02/12/2024 - Audiência de instrução designada para 20/01/2025.',
        'Valor da causa': 280000,
        'Valor do risco': 95000,
    },
    {
        'Status': 'Ativo',
        'Autor/Réu': 'Gabriela Martins Costa',
        'Número Autos': '0004567-12.2020.8.26.0100',
        'Natureza da ação': 'Consumidor',
        'Vara': '5ª Vara Cível',
        'Comarca': 'São Paulo',
        'Data de início': date(2020, 9, 22),
        'Posição da Renault': 'passivo',
        'Resumo do caso': 'Alegação de vício oculto em veículo, com defeito recorrente não sanado em garantia.',
        'Defesa': 'Sim',
        'Fase processual': 'Execução',
        'Movimentações': '08/12/2024 - Penhora online realizada. Valor bloqueado: R$ 12.400,00.',
        'Valor da causa': 45000,
        'Valor do risco': 12000,
    },
    {
        'Status': 'Ativo',
        'Autor/Réu': 'Fazenda Nacional vs Renault Geely Brasil Ltda',
        'Número Autos': '0018901-90.2021.4.13.6000',
        'Natureza da ação': 'Tributário',
        'Vara': '2ª Vara Federal',
        'Comarca': 'Curitiba',
        'Data de início': date(2021, 7, 12),
        'Posição da Renault': 'passivo',
        'Resumo do caso': 'Execução fiscal referente a recolhimento de IPI sobre veículos exportados.',
        'Defesa': 'Sim',
        'Fase processual': 'Execução',
        'Movimentações': '20/12/2024 - Recurso Especial pautado para julgamento no STJ.',
        'Valor da causa': 12500000,
        'Valor do risco': 4800000,
    },
    {
        'Status': 'Ativo',
        'Autor/Réu': 'Auto Peças Dinâmica Ltda',
        'Número Autos': '0008901-90.2023.8.16.0001',
        'Natureza da ação': 'Contratual',
        'Vara': '1ª Vara Empresarial',
        'Comarca': 'Curitiba',
        'Data de início': date(2023, 8, 5),
        'Posição da Renault': 'passivo',
        'Resumo do caso': 'Fornecedor requer indenização por rescisão de contrato de fornecimento exclusivo.',
        'Defesa': 'Não',
        'Fase processual': 'Conhecimento',
        'Movimentações': '28/11/2024 - Prazo de defesa prorrogado por 30 dias.',
        'Valor da causa': 620000,
        'Valor do risco': 190000,
    },
    {
        'Status': 'Ativo',
        'Autor/Réu': 'Distribuidora Paulista de Automóveis SA',
        'Número Autos': '0009012-12.2021.8.26.0001',
        'Natureza da ação': 'Contratual',
        'Vara': '2ª Vara Empresarial',
        'Comarca': 'São Paulo',
        'Data de início': date(2021, 11, 22),
        'Posição da Renault': 'ativo',
        'Resumo do caso': 'Ação de cobrança contra ex-concessionária com passivo de estoque pendente.',
        'Defesa': 'Sim',
        'Fase processual': 'Cumprimento de sentença',
        'Movimentações': '05/12/2024 - Expedição de mandado de penhora e avaliação de bens.',
        'Valor da causa': 380000,
        'Valor do risco': 45000,
    },
    {
        'Status': 'Inativo',
        'Autor/Réu': 'Roberto Santos Oliveira',
        'Número Autos': '0011234-56.2019.5.09.0015',
        'Natureza da ação': 'Trabalhista',
        'Vara': '3ª Vara do Trabalho',
        'Comarca': 'Curitiba',
        'Data de início': date(2019, 5, 10),
        'Posição da Renault': 'passivo',
        'Resumo do caso': 'Pedido de horas extras e adicional de periculosidade. Encerrado por acordo.',
        'Defesa': 'Sim',
        'Fase processual': 'Encerrado',
        'Movimentações': '15/08/2022 - Acordo homologado. Processo encerrado.',
        'Valor da causa': 95000,
        'Valor do risco': 0,
    },
]


def construir(destino: Path) -> Path:
    wb = Workbook()

    # ------------------------------------------------------------------
    # Aba 1 — Processos (é esta que o sistema lê)
    # ------------------------------------------------------------------
    ws = wb.active
    ws.title = 'Processos'

    borda = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9'),
    )

    cabecalhos = [coluna[0] for coluna in COLUNAS]

    # IMPORTANTE: o cabeçalho tem de ficar na linha 1 — o importador lê a
    # primeira linha da primeira aba como cabeçalho.
    for indice, (titulo, largura, obrigatoria, _fmt, _ajuda, _alias) in enumerate(COLUNAS, start=1):
        celula = ws.cell(row=1, column=indice, value=titulo + (' *' if obrigatoria else ''))
        celula.font = Font(name=FONTE, size=10, bold=True, color='FFFFFF')
        celula.fill = PatternFill('solid', fgColor=AZUL_ESCURO if not obrigatoria else AZUL_RENAULT)
        celula.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        celula.border = borda
        ws.column_dimensions[get_column_letter(indice)].width = largura

    ws.row_dimensions[1].height = 30

    # Linhas de exemplo: fundo amarelo para deixar claro que devem ser apagadas.
    for linha_indice, exemplo in enumerate(EXEMPLOS, start=2):
        for coluna_indice, (titulo, _l, _o, formato, _a, _al) in enumerate(COLUNAS, start=1):
            celula = ws.cell(row=linha_indice, column=coluna_indice, value=exemplo.get(titulo))
            celula.font = Font(name=FONTE, size=10)
            celula.alignment = Alignment(vertical='top', wrap_text=titulo in {'Resumo do caso', 'Movimentações'})
            celula.border = borda
            celula.fill = PatternFill('solid', fgColor=AMARELO)
            if formato:
                celula.number_format = formato
        ws.row_dimensions[linha_indice].height = 30

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLUNAS))}{len(EXEMPLOS) + 1}'

    # Listas suspensas nas colunas de domínio fechado, até a linha 1000.
    validacoes = {
        'Status': '"Ativo,Inativo"',
        'Natureza da ação': '"Trabalhista,Cível,Consumidor,Tributário,Contratual,Administrativo,Outros"',
        'Posição da Renault': '"ativo,passivo"',
        'Defesa': '"Sim,Não"',
        'Fase processual': '"Conhecimento,Recurso,Execução,Cumprimento de sentença,Encerrado,Outros"',
    }
    for titulo, formula in validacoes.items():
        letra = get_column_letter(cabecalhos.index(titulo) + 1)
        validacao = DataValidation(type='list', formula1=formula, allow_blank=True, showDropDown=False)
        validacao.error = 'Valor fora da lista. O sistema ainda aceita variações, mas prefira estas opções.'
        validacao.errorTitle = 'Valor sugerido'
        validacao.errorStyle = 'warning'
        ws.add_data_validation(validacao)
        validacao.add(f'{letra}2:{letra}1000')

    # As instruções vão em comentário no cabeçalho, não em uma linha da planilha:
    # qualquer texto solto abaixo dos dados seria lido como um processo inválido
    # na importação.
    instrucoes = Comment(
        'COMO PREENCHER\n\n'
        f'1. Apague as linhas amarelas de exemplo (2 a {len(EXEMPLOS) + 1}) e preencha a partir da linha 2.\n'
        '2. Não mova este cabeçalho nem insira título/logo acima dele — o sistema lê a linha 1 como cabeçalho.\n'
        '3. Colunas com * são obrigatórias: sem elas a linha é descartada.\n'
        '4. Não acrescente texto solto abaixo dos dados: vira uma linha inválida na importação.\n\n'
        'Detalhes de cada coluna na aba "Como preencher".',
        'Sistema de Gestão Jurídica',
    )
    instrucoes.width = 380
    instrucoes.height = 200
    ws['A1'].comment = instrucoes

    # ------------------------------------------------------------------
    # Aba 2 — Como preencher
    # ------------------------------------------------------------------
    guia = wb.create_sheet('Como preencher')

    titulo = guia.cell(row=1, column=1, value='Renault Geely — Gestão Jurídica · Como preencher a planilha')
    titulo.font = Font(name=FONTE, size=13, bold=True, color=AZUL_ESCURO)
    guia.merge_cells('A1:E1')
    guia.row_dimensions[1].height = 24

    regras = [
        'O cabeçalho precisa estar na LINHA 1 da primeira aba. Não insira título, logo ou linhas em branco acima.',
        'Não use células mescladas no cabeçalho.',
        'Só "Autor/Réu" e "Número Autos" são obrigatórios — sem eles a linha é descartada na importação.',
        'As demais colunas podem ficar em branco: o processo entra assim mesmo, com um aviso na tela de inconsistências.',
        'Os nomes das colunas não precisam ser idênticos: a coluna "Nomes também aceitos" mostra as variações reconhecidas.',
        'Reenviar um processo com o mesmo Número Autos atualiza o registro — não duplica.',
        'Apague as linhas amarelas de exemplo antes de enviar: elas são fictícias e entrariam no sistema.',
        'Não deixe texto solto abaixo dos dados (observações, totais, rodapé): vira uma linha inválida.',
        'Formatos aceitos do arquivo: XLSX, XLS ou CSV, até 25 MB.',
    ]
    for indice, regra in enumerate(regras, start=3):
        celula = guia.cell(row=indice, column=1, value='•  ' + regra)
        celula.font = Font(name=FONTE, size=10)
        celula.alignment = Alignment(vertical='center', wrap_text=True)
        guia.merge_cells(start_row=indice, start_column=1, end_row=indice, end_column=5)
        guia.row_dimensions[indice].height = 16

    inicio_tabela = len(regras) + 4
    for indice, texto in enumerate(
        ['Coluna', 'Obrigatória', 'O que preencher', 'Exemplo', 'Nomes também aceitos'], start=1
    ):
        celula = guia.cell(row=inicio_tabela, column=indice, value=texto)
        celula.font = Font(name=FONTE, size=10, bold=True, color='FFFFFF')
        celula.fill = PatternFill('solid', fgColor=AZUL_RENAULT)
        celula.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    primeiro_exemplo = EXEMPLOS[0]
    for deslocamento, (titulo_coluna, _l, obrigatoria, _f, ajuda, alias) in enumerate(COLUNAS, start=1):
        linha = inicio_tabela + deslocamento
        valor_exemplo = primeiro_exemplo.get(titulo_coluna)
        if isinstance(valor_exemplo, date):
            valor_exemplo = valor_exemplo.strftime('%d/%m/%Y')
        elif isinstance(valor_exemplo, (int, float)):
            valor_exemplo = f'R$ {valor_exemplo:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
        elif isinstance(valor_exemplo, str):
            valor_exemplo = valor_exemplo.split('\n')[0][:70]

        valores = [
            titulo_coluna,
            'Sim' if obrigatoria else 'Não',
            ajuda,
            valor_exemplo,
            alias,
        ]
        for coluna_indice, valor in enumerate(valores, start=1):
            celula = guia.cell(row=linha, column=coluna_indice, value=valor)
            celula.font = Font(name=FONTE, size=10, bold=coluna_indice == 1)
            celula.alignment = Alignment(vertical='top', wrap_text=True)
            if obrigatoria:
                celula.fill = PatternFill('solid', fgColor=AMARELO)
            elif deslocamento % 2 == 0:
                celula.fill = PatternFill('solid', fgColor=CINZA_CLARO)
        guia.row_dimensions[linha].height = 42

    for letra, largura in zip('ABCDE', [22, 12, 62, 34, 42]):
        guia.column_dimensions[letra].width = largura

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


if __name__ == '__main__':
    saida = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[2] / 'docs' / 'modelo-planilha-importacao.xlsx'
    print(construir(saida))
