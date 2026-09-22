"""Leitura das fórmulas da planilha financeira.

Duas necessidades:

1. **Critério do realizado.** Nas abas RAP e Resultado, cada linha soma os
   lançamentos com um SUMIF/SUMIFS diferente — e nem sempre o critério é o
   óbvio (a linha "CORPORATE / Consumidor - Condenação" soma ÁREA = "consumidor";
   a linha "TRIBUTÁRIO" soma o EDOA "CORPORATE"). Em vez de adivinhar, lemos a
   fórmula e extraímos o critério, resolvendo as referências de célula.

2. **Valor de célula com fórmula sem cache.** Um arquivo gerado por script
   (sem ter passado pelo Excel) não guarda o resultado das fórmulas. Para
   contas simples como "=7457.5+8011.26" calculamos o valor com segurança.
"""

from __future__ import annotations

import ast
import io
import operator
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

# Coluna da aba "Controle Lançamentos" -> campo do lançamento.
# Letras conforme o layout da planilha original (A=STATUS ... U=PAGO).
COLUNAS_LANCAMENTO_POR_LETRA = {
    "B": "categoria",
    "C": "area",
    "D": "tipo_pagamento",
    "F": "motivo",
    "H": "mes_referencia",
    "J": "recorrencia",
    "Q": "edoa",
    "R": "centro_custo",
    "S": "conta_contabil",
}


@dataclass
class Grade:
    """Uma aba lida em memória: valores (resultado em cache) e fórmulas."""

    nome: str
    valores: list[list[object]]
    formulas: list[list[object]]

    def valor(self, linha: int, coluna: int) -> object:
        """Linha e coluna 1-based, como no Excel."""
        if 1 <= linha <= len(self.valores) and 1 <= coluna <= len(self.valores[linha - 1]):
            return self.valores[linha - 1][coluna - 1]
        return None

    def formula(self, linha: int, coluna: int) -> object:
        if 1 <= linha <= len(self.formulas) and 1 <= coluna <= len(self.formulas[linha - 1]):
            return self.formulas[linha - 1][coluna - 1]
        return None

    @property
    def total_linhas(self) -> int:
        return len(self.valores)


def carregar_grades(conteudo: bytes, abas: list[str]) -> dict[str, Grade]:
    """Lê só as abas pedidas, em modo leitura (rápido mesmo em arquivo grande)."""
    grades: dict[str, Grade] = {}
    wb_valores = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    wb_formulas = load_workbook(io.BytesIO(conteudo), data_only=False, read_only=True)
    try:
        for aba in abas:
            if aba not in wb_valores.sheetnames:
                continue
            valores = [list(linha) for linha in wb_valores[aba].iter_rows(values_only=True)]
            formulas = [list(linha) for linha in wb_formulas[aba].iter_rows(values_only=True)]
            grades[aba] = Grade(nome=aba, valores=valores, formulas=formulas)
    finally:
        wb_valores.close()
        wb_formulas.close()
    return grades


# --------------------------------------------------------------------------
# SUMIF / SUMIFS
# --------------------------------------------------------------------------

_REF_CELULA = re.compile(r"^(?:'(?P<aba>[^']*)'!|(?P<aba2>[\w]+)!)?\$?(?P<col>[A-Z]{1,3})\$?(?P<lin>\d+)$")


@dataclass
class Criterio:
    """Critério de soma extraído de uma fórmula."""

    filtros: dict[str, str] = field(default_factory=dict)  # campo -> valor
    valido: bool = True
    observacao: str | None = None


def _dividir_argumentos(corpo: str) -> list[str]:
    """Separa argumentos por vírgula, respeitando aspas e parênteses."""
    partes, atual, profundidade, em_aspas = [], [], 0, False
    for caractere in corpo:
        if caractere == '"':
            em_aspas = not em_aspas
        if not em_aspas:
            if caractere == "(":
                profundidade += 1
            elif caractere == ")":
                profundidade -= 1
            elif caractere == "," and profundidade == 0:
                partes.append("".join(atual).strip())
                atual = []
                continue
        atual.append(caractere)
    if atual:
        partes.append("".join(atual).strip())
    return partes


def _letra_da_coluna(argumento: str) -> str | None:
    """"'Controle Lançamentos'!Q:Q" -> "Q"."""
    parte = argumento.split("!")[-1]
    achou = re.fullmatch(r"\$?([A-Z]{1,3}):\$?([A-Z]{1,3})", parte.strip())
    return achou.group(1) if achou else None


def _resolver(argumento: str, grade: Grade) -> tuple[str | None, str | None]:
    """Resolve um critério: literal entre aspas ou referência a célula da mesma aba.

    Devolve (valor, problema).
    """
    argumento = argumento.strip()
    if argumento.startswith('"') and argumento.endswith('"'):
        return argumento[1:-1], None
    if "#REF!" in argumento:
        return None, "referência quebrada (#REF!)"

    achou = _REF_CELULA.match(argumento)
    if not achou:
        return None, f"critério não suportado ({argumento})"

    aba_ref = achou.group("aba") or achou.group("aba2")
    if aba_ref and aba_ref.strip() != grade.nome.strip():
        return None, f"critério aponta para outra aba ({aba_ref})"

    coluna = column_index_from_string(achou.group("col"))
    linha = int(achou.group("lin"))
    valor = grade.valor(linha, coluna)
    if valor is None:
        # A célula pode ser outra fórmula simples ("=$C$2"): resolve um nível.
        formula = grade.formula(linha, coluna)
        if isinstance(formula, str) and formula.startswith("="):
            return _resolver(formula[1:], grade)
        return None, (
            "a fórmula usa como critério uma célula vazia — no Excel esta linha dá sempre zero"
        )
    return str(valor).strip(), None


def extrair_criterio_soma(formula: object, grade: Grade) -> Criterio | None:
    """Lê SUMIF/SUMIFS sobre a aba de lançamentos e devolve o critério.

    Devolve None quando a célula não é uma fórmula de soma (valor digitado ou vazia).
    """
    if not isinstance(formula, str) or not formula.startswith("="):
        return None

    texto = formula[1:].strip()
    achou = re.match(r"^(SUMIFS?|SOMASES?|SOMASE)\((.*)\)$", texto, flags=re.IGNORECASE | re.DOTALL)
    if not achou:
        return Criterio(valido=False, observacao="fórmula que não é SUMIF/SUMIFS")

    nome = achou.group(1).upper()
    argumentos = _dividir_argumentos(achou.group(2))
    criterio = Criterio()

    if nome in {"SUMIF", "SOMASE"}:
        # SUMIF(intervalo_criterio, criterio, intervalo_soma)
        if len(argumentos) < 2:
            return Criterio(valido=False, observacao="SUMIF incompleto")
        pares = [(argumentos[0], argumentos[1])]
    else:
        # SUMIFS(intervalo_soma, intervalo1, criterio1, intervalo2, criterio2, ...)
        pares = list(zip(argumentos[1::2], argumentos[2::2]))

    for intervalo, bruto in pares:
        letra = _letra_da_coluna(intervalo)
        campo = COLUNAS_LANCAMENTO_POR_LETRA.get(letra or "")
        if campo is None:
            criterio.valido = False
            criterio.observacao = f"coluna de critério não reconhecida ({intervalo})"
            continue
        valor, problema = _resolver(bruto, grade)
        if problema:
            criterio.valido = False
            criterio.observacao = problema
            continue
        criterio.filtros[campo] = valor

    return criterio


# --------------------------------------------------------------------------
# Aritmética simples ("=7457.5+8011.26")
# --------------------------------------------------------------------------

_OPERADORES = {
    ast.Add: operator.add,
    ast.Sub: operator.sub,
    ast.Mult: operator.mul,
    ast.Div: operator.truediv,
    ast.USub: operator.neg,
    ast.UAdd: operator.pos,
}


def avaliar_aritmetica(formula: object) -> float | None:
    """Calcula fórmulas que são só números e + - * / ( ). Qualquer outra coisa: None.

    Não usa eval(): percorre a árvore sintática aceitando apenas números e
    operadores aritméticos — uma fórmula maliciosa não executa nada.
    """
    if not isinstance(formula, str) or not formula.startswith("="):
        return None
    expressao = formula[1:].strip().replace(",", ".")
    if not re.fullmatch(r"[\d.\s+\-*/()]+", expressao):
        return None
    try:
        arvore = ast.parse(expressao, mode="eval")
    except SyntaxError:
        return None

    def calcular(no):
        if isinstance(no, ast.Expression):
            return calcular(no.body)
        if isinstance(no, ast.Constant) and isinstance(no.value, (int, float)):
            return float(no.value)
        if isinstance(no, ast.BinOp) and type(no.op) in _OPERADORES:
            return _OPERADORES[type(no.op)](calcular(no.left), calcular(no.right))
        if isinstance(no, ast.UnaryOp) and type(no.op) in _OPERADORES:
            return _OPERADORES[type(no.op)](calcular(no.operand))
        raise ValueError("expressão não permitida")

    try:
        return round(calcular(arvore), 2)
    except (ValueError, ZeroDivisionError):
        return None
