"""Gera uma planilha financeira de exemplo, com o MESMO layout da planilha do Jurídico.

Mesmas abas, mesmos cabeçalhos nas mesmas posições e as mesmas fórmulas
(SUMIF/SUMIFS do RAP e do Resultado, IF do STATUS) — mas com dados fictícios:
nenhum fornecedor, pessoa ou valor real. Serve para:

- os testes automatizados (não dependem do arquivo real, que tem dados sensíveis);
- o botão "Baixar planilha modelo" da tela de importação.

Os defeitos que a planilha real tem foram reproduzidos de propósito, para os
testes garantirem que o sistema os detecta: linha do RAP cujo critério é uma
célula vazia, duas linhas somando o mesmo critério, REAL VARIÁVEL procurando
"Variável" e linha "Juros e correção" sem fórmula no Resultado.
"""

from __future__ import annotations

import io
import random

from openpyxl import Workbook

MESES = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]

CL = "'Controle Lançamentos'"
RAP = "'RAP PAGAMENTOS'"

ESCRITORIOS = [
    "Escritório Alfa Advogados",
    "Beta & Associados",
    "Gama Consultoria Jurídica",
    "Delta Advocacia",
]

# (categoria, área, EDOA, motivo, recorrência, valor base)
PERFIS = [
    ("RDB - CONS", "Consumidor", "HONORÁRIOS E DESPESAS", "Honorários", "Fixo", 30000),
    ("RDB - CONS", "Consumidor", "CORPORATE", "Condenação", "Variável pontual", 45000),
    ("RDB - CONS", "Consumidor", "ACORDOS CONSUMIDOR", "Acordo", "Variável pontual", 12000),
    ("RDB - CONS", "Acordo Projeto SAC", "ACORDOS SAC", "Acordo", "Variável pontual", 8000),
    ("RDB - DIV", "Cível", "HONORÁRIOS E DESPESAS", "Honorários", "Fixo", 6400),
    ("RDB - DIV", "Cível", "HONORÁRIOS E DESPESAS", "Custas", "Variável pontual", 900),
    ("RDB - TRAB CVP", "Trabalhista", "TRABALHISTA GERAL", "Condenação", "Variável pontual", 25000),
    ("RDB - TRAB CVP", "Trabalhista", "TRABALHISTA GERAL", "acordo", "variável pontual", 15000),
    ("RDB - TRAB SUPORTE", "Trabalhista", "TRABALHISTA ACORDOS", "Acordo", "Variável pontual", 9000),
    ("RDB - TRAB PIF", "Trabalhista", "PIF/PDV", "Acordo", "Variável pontual", 20000),
    ("RDB - TRIB", "Tributário", "HONORÁRIOS E DESPESAS", "Honorários", "Fixo", 17000),
    ("RDB - TRIB", "Tributário", "CORPORATE", "Custas", "Variável pontual", 3000),
    ("RDB - DIV", "LGPD ", "LGPD GERAL & TERCEIROS", "Honorários", "Fixo", 12000),
    ("RDB - DIV", "LGPD ", "LGPD GERAL & TERCEIROS", "LGPD ", "Variável pontual", 2500),
    ("RDB - DIV", "Cartório", "CARTÓRIO", "Despesas", "Fixo valor variável ", 1500),
    ("RDB - DIV", "Projuris SUPRAMONTE", "SISTEMAS", "Projuris", "Fixo", 2068.72),
    ("RDB - DIV", "Consumidor", "OUTRA ÁREA", "Juros e correção", "Variável pontual", 700),
]


def _status_formula(linha: int) -> str:
    return (
        f'=IF(U{linha}<>"","Pago",IF(T{linha}<>"","Enviado para pagamento",'
        f'IF(N{linha}<>"","Pedido Recepcionado",IF(M{linha}<>"","Pedido Concluído",'
        f'IF(L{linha}<>"","RC Criada","Lançar Pgto")))))'
    )


def _aba_base(wb: Workbook) -> None:
    ws = wb.create_sheet("Base")
    colunas = {
        "Área": ["Acordo Projeto SAC", "Cartório", "Cível", "Consumidor", "LGPD ", "Trabalhista", "Tributário", "Societário", "Rede", "Serviços terceirizados", "Projuris SUPRAMONTE"],
        "PAGAMENTO": ["Adiantamento", "Interno"],
        "Fornecedor": ESCRITORIOS,
        "Motivo": ["Acordo", "Projuris", "Condenação", "Custas", "Depósito Judicial", "Despesas", "Honorários", "Juros e correção", "LGPD"],
        "Fixo/Varíavel": ["Fixo", "Fixo valor variável ", "Variável pontual"],
        "EDOA": ["CARTÓRIO", "Depósito Judicial ", "HONORÁRIOS E DESPESAS", "LGPD GERAL & TERCEIROS", "OUTRA ÁREA", "OUTROS (Jurídico)", "CORPORATE", "ACORDOS CONSUMIDOR", "ACORDOS SAC", "TRABALHISTA GERAL", "TRABALHISTA ACORDOS", "PIF/PDV", "SISTEMAS"],
    }
    for c, (titulo, valores) in enumerate(colunas.items(), start=1):
        ws.cell(1, c, titulo)
        for l, valor in enumerate(valores, start=2):
            ws.cell(l, c, valor)


def _aba_lancamentos(wb: Workbook, ano: int, gerador: random.Random) -> list[dict]:
    ws = wb.create_sheet("Controle Lançamentos")
    cabecalho = [
        "STATUS", "CATEGORIA", "ÁREA", "PAGAMENTO", "REFERÊNCIA", "MOTIVO", "Valor", "MÊS REFERENCIA",
        "DATA PAGAMENTO", "RECORRÊNCIA", "DESCRIÇÃO", "RC", "Nº PEDIDO", "Recepção", "ITEM", "CHAMADO",
        "EDOA", "CENTRO DE CUSTO", "CONTA CONTÁBIL", "ENVIO PARA PAGAMENTO", "PAGO / Nº DOCUMENTO",
    ]
    for c, titulo in enumerate(cabecalho, start=1):
        ws.cell(1, c, titulo)

    gerados = []
    linha = 2
    referencia = 1000
    for mes in range(1, 7):
        for perfil in PERFIS:
            categoria, area, edoa, motivo, recorrencia, base = perfil
            referencia += 1
            valor = round(base * gerador.uniform(0.8, 1.2), 2) if recorrencia.strip().lower() != "fixo" else base
            # Etapa do fluxo: meses antigos já pagos; os recentes em andamento.
            etapa = 5 if mes <= 3 else gerador.randint(0, 5)
            dados = {
                "B": categoria, "C": area, "D": gerador.choice(["Interno", "ADIANTAMENTO"]), "E": referencia,
                "F": motivo, "G": valor, "H": MESES[mes - 1],
                "I": f"{gerador.randint(5, 28):02d}.{min(mes + 1, 12):02d}.{ano}", "J": recorrencia,
                "K": f"{motivo.strip()} {area.strip()} {MESES[mes - 1].lower()}",
                "L": 25000000 + referencia if etapa >= 1 else None,
                "M": 4500000000 + referencia if etapa >= 2 else None,
                "N": 180000000 + referencia if etapa >= 3 else None,
                "O": 10, "P": f"{1300000 + referencia}\xa0" if etapa >= 1 else None,
                "Q": edoa, "R": " GI52000" if mes % 2 else "GI52015", "S": 622600,
                "T": f"{gerador.randint(1, 28):02d}.{mes:02d}.{ano}" if etapa >= 4 else None,
                "U": 211400000 + referencia if etapa >= 5 else None,
            }
            ws.cell(linha, 1, _status_formula(linha))
            for letra, valor_celula in dados.items():
                if valor_celula is not None:
                    ws[f"{letra}{linha}"] = valor_celula
            gerados.append({**dados, "linha": linha, "mes": mes})
            linha += 1
    return gerados


def _aba_rap(wb: Workbook, ano: int) -> None:
    ws = wb.create_sheet("RAP PAGAMENTOS")
    ws["B2"], ws["C2"], ws["I2"] = "MÊS FECHAMENTO", 2, "consumidor"
    ws["B3"] = "RESTOS A PAGAR POR MÊS"
    for c, titulo in zip("BCDEFGH", ["EDOA", "ÁREA", "BUDGET \nANUAL", "PREVISTO", "REALIZADO", "GAP", "BUDGET RESTANTE"]):
        ws[f"{c}4"] = titulo

    def sumif(l: int) -> str:
        return f"=SUMIF({CL}!Q:Q,{RAP}!B{l},{CL}!G:G)"

    def sumifs(l: int, area_ref: str, edoa_ref: str) -> str:
        return f"=SUMIFS({CL}!G:G,{CL}!C:C,{area_ref},{CL}!Q:Q,{edoa_ref})"

    # (EDOA, área, budget, fórmula do realizado)
    linhas = [
        ("OUTROS (Jurídico)", None, 60000, "edoa"),
        ("CARTÓRIO", None, 140000, "edoa"),
        ("SISTEMAS", None, 60000, 3),  # valor digitado no lugar da fórmula (como na planilha real)
        ("SISTEMAS", "Projuris SUPRAMONTE", 30000, ("C", "$B$7")),
        ("HONORÁRIOS E DESPESAS", None, 1500000, "edoa"),
        ("HONORÁRIOS E DESPESAS", "Cível", 90000, ("C", "$B$9")),
        ("HONORÁRIOS E DESPESAS", "Consumidor", 400000, ("C", "$B$9")),
        ("HONORÁRIOS E DESPESAS", "Tributário", 210000, ("C", "$B$9")),
        ("PIF/PDV", None, 150000, "edoa"),
        ("CORPORATE", None, 700000, "edoa"),
        # Área digitada na fórmula: soma ÁREA "consumidor" do EDOA CORPORATE.
        ("CORPORATE", "Consumidor - Condenação", 600000, '"consumidor"'),
        ("ACORDOS CONSUMIDOR", "Consumidor - Acordo", 100000, ("I2", "$B$16")),
        ("ACORDOS SAC", "Consumidor - Projeto SAC", 60000, ('"Acordo Projeto SAC"', "$B$17")),
        # Linha de outro EDOA somando o CORPORATE — mesmo defeito do "TRIBUTÁRIO" real.
        ("TRIBUTÁRIO", "Tributário", 20000, ("C", "$B$14")),
        # Critério em célula vazia (C19) — no Excel dá sempre zero.
        ("DEPÓSITO JUDICIAL", None, 0, ("C", "$B$14")),
        # Mesmo critério da linha 15 (CORPORATE/consumidor): o gasto aparece duas vezes.
        ("DEPÓSITO JUDICIAL", "Consumidor", 0, ("C", "$B$14")),
        ("LGPD GERAL & TERCEIROS", "LGPD", 200000, "edoa"),
        ("TRABALHISTA ACORDOS", "Trabalhista", 80000, "edoa"),
        ("TRABALHISTA GERAL", "Trabalhista", 1100000, "edoa"),
    ]
    for l, (edoa, area, budget, regra) in enumerate(linhas, start=5):
        ws[f"B{l}"] = edoa
        if area:
            ws[f"C{l}"] = area
        ws[f"D{l}"] = budget
        ws[f"E{l}"] = f"=(D{l}/12)*$C$2"
        if regra == "edoa":
            ws[f"F{l}"] = sumif(l)
        elif isinstance(regra, (int, float)):
            ws[f"F{l}"] = regra
        elif isinstance(regra, str):
            ws[f"F{l}"] = sumifs(l, regra, f"{RAP}!$B$14")
        else:
            area_ref, edoa_ref = regra
            area_ref = f"{RAP}!{area_ref}{l}" if area_ref == "C" else (area_ref if area_ref.startswith('"') else f"{RAP}!{area_ref}")
            ws[f"F{l}"] = sumifs(l, area_ref, f"{RAP}!{edoa_ref}")
        ws[f"G{l}"] = f"=E{l}-F{l}"
        ws[f"H{l}"] = f"=D{l}-F{l}"

    # Bloco 2 — honorários por área, fixo × variável
    ws["B30"] = "HONORÁRIOS E DESPESAS"
    titulos = {
        "B": "ÁREA", "D": "BUDGET \nANUAL", "E": "PREVISTO", "F": "REAL \nTOTAL", "G": "GAP", "H": "BUDGET RESTANTE",
        "I": "BUDGET\nVARIÁVEL", "J": "PREVISTO\nVARIÁVEL", "K": "REAL\nVARIÁVEL", "L": "GAP\nVARIÁVEL",
        "M": "RESTANTE VARIÁVEL ATÉ DEZ", "N": "BUDGET \nFIXO", "O": "PREVISTO\nFIXO", "P": "REAL\nFIXO",
        "Q": "GAP\nFIXO", "R": "RESTANTE FIXO ATÉ DEZ",
    }
    for c, titulo in titulos.items():
        ws[f"{c}31"] = titulo
    for l, (area, total, variavel, fixo) in enumerate(
        [("Cível", 90000, 20000, 70000), ("Consumidor", 400000, 40000, 360000), ("Tributário", 210000, 6000, 204000)],
        start=32,
    ):
        ws[f"B{l}"], ws[f"C{l}"] = area, area
        ws[f"D{l}"], ws[f"I{l}"], ws[f"N{l}"] = total, variavel, fixo
        base = f"{CL}!G:G,{CL}!C:C,{RAP}!B{l},{CL}!Q:Q,{RAP}!$B$9"
        ws[f"E{l}"] = f"=(D{l}/12)*$C$2"
        ws[f"F{l}"] = f"=SUMIFS({base})"
        ws[f"G{l}"], ws[f"H{l}"] = f"=E{l}-F{l}", f"=D{l}-F{l}"
        ws[f"J{l}"] = f"=(I{l}/12)*$C$2"
        ws[f"K{l}"] = f'=SUMIFS({base},{CL}!J:J,"Variável")'
        ws[f"L{l}"], ws[f"M{l}"] = f"=J{l}-K{l}", f"=I{l}-K{l}"
        ws[f"O{l}"] = f"=(N{l}/12)*$C$2"
        ws[f"P{l}"] = f'=SUMIFS({base},{CL}!J:J,"Fixo")'
        ws[f"Q{l}"], ws[f"R{l}"] = f"=O{l}-P{l}", f"=N{l}-P{l}"


def _aba_resultado(wb: Workbook) -> None:
    ws = wb.create_sheet("Resultado")
    ws["B1"], ws["B2"], ws["C2"] = "APENAS UM FILTRO POR MESES", "MÊS", "Março"
    for c, titulo in zip("BCDEFG", ["CATEGORIA", "DOA", "PERÍODO", "IMPACTO", "R$ GASTO MENSAL", "R$ GASTO ANUAL"]):
        ws[f"{c}4"] = titulo
    linhas = [
        ("RDB - CONS", "ACORDOS CONSUMIDOR + SAC", "APCO", False),  # sem fórmula (como na real)
        ("RDB - CONS", "CORPORATE", "APCO", True),
        ("RDB - CONS", "HONORÁRIOS E DESPESAS", "G&A", True),
        ("RDB - TRAB CVP", "TRABALHISTA GERAL", "MASSA", True),
        ("RDB - TRAB SUPORTE", "TRABALHISTA ACORDOS", "G&A", True),
        ("RDB - TRAB PIF", "PIF/PDV", "APCE", True),
        ("RDB - TRIB", "CORPORATE", "APCO", True),
        ("Juros e correção", None, "Juros e correção", False),
    ]
    for l, (categoria, doa, impacto, com_formula) in enumerate(linhas, start=5):
        ws[f"B{l}"] = categoria
        if doa:
            ws[f"C{l}"] = doa
        ws[f"D{l}"] = "=$C$2"
        ws[f"E{l}"] = impacto
        if com_formula:
            ws[f"F{l}"] = f"=SUMIFS({CL}!G:G,{CL}!B:B,B{l},{CL}!H:H,D{l},{CL}!Q:Q,Resultado!C{l})"
            ws[f"G{l}"] = f"=SUMIFS({CL}!G:G,{CL}!B:B,B{l},{CL}!Q:Q,Resultado!C{l})"
    ws["G14"] = "=SUM(G6:G11)"


def _aba_lgpd(wb: Workbook) -> None:
    ws = wb.create_sheet("LGPD")
    ws["A2"], ws["B2"] = "EDOA", "LGPD GERAL & TERCEIROS"
    ws["A4"], ws["B4"] = "Soma de Valor", "Rótulos de Coluna"
    for c, titulo in zip("ABCD", ["Rótulos de Linha", "Fixo", "Variável pontual", "Total Geral"]):
        ws[f"{c}5"] = titulo


def _aba_adiantamentos(wb: Workbook) -> None:
    ws = wb.create_sheet("Adiantamentos")
    cabecalho = [
        "Status", "Ok na Planilha de pagamentos", "Área", "Valor", "Chamado do adiantamento", "DOA",
        "Nº chamado da baixa", "Valor Excedente", "Valor total do adiantamento", "Baixa do Adiantamento",
        None, "Escritório",
    ]
    for c, titulo in enumerate(cabecalho, start=1):
        if titulo:
            ws.cell(1, c, titulo)
    registros = [
        ("Ok", "Trabalhista", 500000, 1317001, "TRABALHISTA", 1345001, 1200.5, 294918001, ESCRITORIOS[0]),
        ("Ok", "Consumidor", 300000, 1322002, "CORPORATE", 1343002, 0, None, ESCRITORIOS[1]),
        (None, "Consumidor", 250000, 1399003, "CORPORATE", None, 0, None, ESCRITORIOS[2]),
        (None, "Trabalhista", 150000, None, "TRABALHISTA", None, 0, None, ESCRITORIOS[3]),
    ]
    for l, (ok, area, valor, chamado, doa, baixa, excedente, documento, escritorio) in enumerate(registros, start=2):
        ws[f"A{l}"] = f'=IF(J{l}<>"","Adiantamento Baixado",IF(G{l}<>"","Baixa lançada",IF(E{l}<>"","Adiantamento Lançado","Lançar adiantamento")))'
        for letra, v in zip("BCDEFGHJL", [ok, area, valor, chamado, doa, baixa, excedente, documento, escritorio]):
            if v is not None:
                ws[f"{letra}{l}"] = v
        ws[f"I{l}"] = f"=D{l}+H{l}"


def _aba_honorarios(wb: Workbook) -> None:
    ws = wb.create_sheet("Honorarios Variaveis ")
    cabecalho = ["DOA", "ÁREA", "DETALHAMENTO", "RECORRÊNCIA", "JAN", "FEV", "MAR", "ABR", "MAI", "JUN",
                 "JUL", "AGO", "SET", "OUT", "NOV", "DEZ", "BUDGET", "REAL", "SALDO"]
    for c, titulo in enumerate(cabecalho, start=1):
        ws.cell(2, c, titulo)
    itens = [
        ("Tributário", "Honorários tributário fixo mensal", "Fixo", [17000] * 12),
        ("Consumidor", "Acompanhamento carteira consumidor", "Fixo", [30000] * 12),
        ("Consumidor", "Mão de obra terceirizada", "Fixo", ["=2500.5"] * 12),  # valor como fórmula
        ("Cível", "Honorários cível fixo mensal", "Fixo", [6400] * 12),
        ("Cível", "Êxito cível", "Variável pontual", [0, 0, 0, 5000, 5000, 5000, 0, 0, 0, 5000, 0, 0]),
    ]
    for l, (area, detalhe, recorrencia, meses) in enumerate(itens, start=3):
        ws.cell(l, 1, "HONORÁRIOS E DESPESAS")
        ws.cell(l, 2, area)
        ws.cell(l, 3, detalhe)
        ws.cell(l, 4, recorrencia)
        for m, valor in enumerate(meses):
            ws.cell(l, 5 + m, valor)
        total = sum(2500.5 if isinstance(v, str) else v for v in meses)
        ws.cell(l, 17, total)
        # Coluna REAL quebrada, como na planilha real.
        ws.cell(l, 18, f"=SUMIFS({CL}!G:G,{CL}!#REF!,'Honorarios Variaveis '!#REF!)")
        ws.cell(l, 19, f"=Q{l}-R{l}")


def _aba_mensais(wb: Workbook) -> None:
    ws = wb.create_sheet("Mensais")
    for c, titulo in enumerate(["Àrea", "Recorrência", "Valor Mensal", "Anual", "TRIMESTRAL", "MÉDIA VARIAVEIS",
                                "Área", "Escritório", "Valor", None, "Area", "Valor", "Anual"], start=1):
        if titulo:
            ws.cell(1, c, titulo)
    contratos = [("Consumidor", "Fixo", 30000, 360000), ("Consumidor", "Fixo", 30000, 360000),
                 ("Tributário", "Fixo", 17000, 204000), ("Cível", "Fixo", 6400, 76800)]
    for l, (area, rec, mensal, anual) in enumerate(contratos, start=2):
        ws.cell(l, 1, area)
        ws.cell(l, 2, rec)
        ws.cell(l, 3, mensal)
        ws.cell(l, 4, anual)
    sistemas = [("Projuris", 2068.72, 24824.64, "Mensal sistema jurídico"), ("Implantação", 5000, 5000, "Parcela única")]
    for l, (nome, valor, anual, obs) in enumerate(sistemas, start=2):
        ws.cell(l, 11, nome)
        ws.cell(l, 12, valor)
        ws.cell(l, 13, anual)
        ws.cell(l, 14, obs)


def _aba_devolucoes(wb: Workbook, ano: int) -> None:
    from datetime import datetime

    ws = wb.create_sheet("Devoluções")
    for c, titulo in enumerate(["ÁREA", "FORNECEDOR", "CÓDIGO DO BANCO", "DATA TRANSF/", "PASTA BENNER",
                                "VALOR DEVOLVIDO", "Nº CHAMADO", "DOCUMENTO", "OBSERVAÇÃO"], start=1):
        ws.cell(1, c, titulo)
    registros = [
        ("Consumidor", ESCRITORIOS[0], 556159, datetime(ano, 1, 24), "Projuris 101", 12000.4, 1339001, 294905001, "RECLASSIFICAR PARA DOA"),
        ("Trabalhista", ESCRITORIOS[1], 556159, datetime(ano, 2, 12), "Projuris 102", 5300, 1341002, 294906002, None),
    ]
    for l, registro in enumerate(registros, start=2):
        for c, v in enumerate(registro, start=1):
            if v is not None:
                ws.cell(l, c, v)
    ws["L2"], ws["M2"] = "Consumidor", "=SUMIF(A:A,L2,F:F)"


def _aba_rf_mensal(wb: Workbook) -> None:
    ws = wb.create_sheet("RF MENSAL")
    contas = [
        ("Jurídico", "FADM", "GI52000", 635110, "Despesas com Estudos e Honorários"),
        ("Jurídico", "FADM", "GI52000", 622600, "Honorários advocatícios"),
        ("Jurídico", "FADM", "GI52015", 611205, "Acordos trabalhistas"),
    ]
    for l, conta in enumerate(contas, start=3):
        for c, v in enumerate(conta, start=1):
            ws.cell(l, c, v)


def gerar_planilha_financeira(ano: int = 2024, semente: int = 42) -> tuple[bytes, list[dict]]:
    """Devolve (conteúdo .xlsx, lançamentos gerados — para os testes conferirem as somas)."""
    gerador = random.Random(semente)
    wb = Workbook()
    wb.remove(wb.active)
    _aba_base(wb)
    _aba_resultado(wb)
    _aba_lgpd(wb)
    lancamentos = _aba_lancamentos(wb, ano, gerador)
    _aba_rap(wb, ano)
    _aba_adiantamentos(wb)
    _aba_honorarios(wb)
    _aba_mensais(wb)
    _aba_devolucoes(wb, ano)
    _aba_rf_mensal(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), lancamentos
